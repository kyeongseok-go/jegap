#!/usr/bin/env python3
"""K-apt 벌크 XLSX → app/data/kapt.json.gz 전처리.
입력: basis.xlsx(기본정보) area.xlsx(면적) fee2023~2026.xlsx(월별 관리비)
출처: k-apt.go.kr 자료실 (주간 갱신). 실행: python3 ingest.py <입력디렉토리>
"""
import sys, json, gzip, warnings, time
from collections import defaultdict
import openpyxl

warnings.filterwarnings("ignore")
SRC = sys.argv[1]
OUT = sys.argv[2] if len(sys.argv) > 2 else "kapt.json.gz"

def rows(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    ws.reset_dimensions()
    it = ws.iter_rows(values_only=True)
    next(it)          # 안내문
    hdr = [str(h) if h else "" for h in next(it)]
    return hdr, it

def num(v):
    if v is None or v == "": return 0.0
    try: return float(v)
    except (TypeError, ValueError): return 0.0

t0 = time.time()

# 1) 기본정보: code → 메타
meta = {}
hdr, it = rows(f"{SRC}/basis.xlsx")
for r in it:
    code = r[4]
    if not code: continue
    used = str(r[11] or "")          # 사용승인일 YYYYMMDD
    if len(used) < 4 or not used[:4].isdigit(): continue
    meta[code] = {
        "name": str(r[5] or "").strip(),
        "sido": str(r[0] or "").strip(),
        "sigungu": str(r[1] or "").strip(),
        "builtYear": int(used[:4]),
        "households": int(num(r[14])),
        "heating": str(r[20] or "").strip() or "정보없음",
        "sale": str(r[10] or "").strip(),      # 분양형태 (분양/임대 — 홈 사례 선정용)
        "kind": str(r[6] or "").strip(),       # 단지분류 (아파트/주상복합 등)
    }
print(f"basis: {len(meta)} danji {time.time()-t0:.0f}s")

# 2) 면적: code → 관리비부과면적
area = {}
hdr, it = rows(f"{SRC}/area.xlsx")
for r in it:
    code, a = r[4], num(r[7])
    if code and a > 0: area[code] = a
print(f"area: {len(area)} {time.time()-t0:.0f}s")

# 3) 관리비 4개년: code → {ym: (total, heating, reserve부과, reserve총적립, 적립률)}
# 컬럼: 4=code 6=YYYYMM 7=공용관리비계 27=개별사용료계 28=난방공용 29=난방전용
#       45=장충금월부과 47=장충금총적립 48=적립률
fees = defaultdict(dict)
for y in (2023, 2024, 2025, 2026):
    hdr, it = rows(f"{SRC}/fee{y}.xlsx")
    assert hdr[4] == "단지코드" and hdr[45] == "장충금 월부과액", f"{y} 컬럼 드리프트: {hdr[4]}, {hdr[45]}"
    n = 0
    for r in it:
        code, ym = r[4], str(r[6] or "")
        if not code or len(ym) != 6: continue
        fees[code][ym] = (
            num(r[7]) + num(r[27]),          # total = 공용관리비 + 개별사용료
            num(r[28]) + num(r[29]),         # heating = 공용 + 전용
            num(r[45]), num(r[47]), num(r[48]),
        )
        n += 1
    print(f"fee{y}: {n} rows {time.time()-t0:.0f}s")

# 4) 조립: 면적으로 ㎡당 환산, 12개월 미만 단지 제외(무소음)
out = []
for code, m in meta.items():
    a = area.get(code)
    if not a or code not in fees: continue
    months = sorted(fees[code].keys())[-40:]      # 최근 40개월
    if len(months) < 12: continue
    pts = []
    for ym in months:
        t, h, rv, rt, rp = fees[code][ym]
        pts.append([ym, round(t / a), round(h / a)])
    last = fees[code][months[-1]]
    out.append({
        "c": code, **m,
        "fees": pts,
        "rv": round(last[2] / a),                 # 장충금 월부과 ㎡당
        "rt": round(last[3]),                     # 총적립액(원)
        "rp": round(last[4], 1),                  # 적립률(%)
    })

print(f"assembled: {len(out)} danji {time.time()-t0:.0f}s")
with gzip.open(OUT, "wt", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
import os
print(f"wrote {OUT}: {os.path.getsize(OUT)/1e6:.1f}MB")
